"""Load user-supplied documents into requirement and table records.

Supported:
- .pdf  ECSS (primary source) and NASA standards, layout-aware (pdfplumber)
- .docx / .doc  ECSS via LibreOffice (tracked changes accepted); other .docx via python-docx
- .txt / .md  generic clause parser (used by the invented test fixtures)

Document identity (doc_id, issue) is read from the Word page header or the cover page (the
standard number and "Rev.N" / "w/Change N" next to it), else from the file name, so downloads keep their names:
`ECSS-E-ST-32C_Rev.115November2008.pdf` -> ECSS-E-ST-32C, Rev.1. An optional sidecar
`<name>.meta.yaml` overrides any field (doc_id, issue, title). A part or volume is part of the doc_id
(`ECSS-E-HB-32-20 Part 1A`, `NASA-HDBK-5010 Vol 1`); a download manifest next to the file is the last resort.
Which files are loaded, and which copy of a document, is the catalogue's job (catalogue.py, ADR-0006).
"""

from __future__ import annotations

import hashlib
import json
import logging
import pickle
import re
from functools import lru_cache
from pathlib import Path

import yaml

from .models import Document
from .parser import is_handbook, parse_paragraphs

SUPPORTED = {".pdf", ".docx", ".doc", ".txt", ".md"}
PREFERENCE = {".pdf": 0, ".docx": 1, ".doc": 2, ".txt": 3, ".md": 3}
CACHE_DIR = ".spacestd-cache"  # parsed documents, next to the documents (gitignored with them)
log = logging.getLogger(__name__)


def _code_version() -> str:
    """Hash of the parsing code: any parser change invalidates the cache."""
    h = hashlib.sha1()
    for name in ("models", "parser", "pdf_extract", "ecss_parser", "nasa_parser", "word_extract", "loaders"):
        h.update((Path(__file__).parent / f"{name}.py").read_bytes())
    return h.hexdigest()[:12]


def read_paragraphs(path: Path) -> list[str]:
    if path.suffix.lower() == ".docx":
        import docx  # python-docx

        d = docx.Document(str(path))
        paragraphs = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return paragraphs
    return path.read_text(encoding="utf-8").splitlines()


# ECSS-E-ST-32-10C, ECSS-Q-ST-70C, NASA-STD-5001B, NASA-HDBK-7005, NASA-STD-10001, NASA-HDBK-8739.19-2 ...
STD_ID_RE = re.compile(
    r"(?<![A-Za-z0-9])(ECSS-[A-Z]-(?:(?:ST|HB|AS|TM)-)?\d{2}(?:-\d{2})*[A-Z]?"
    r"|NASA-(?:STD|HDBK|SPEC)-\d{4,5}(?:\.\d+(?:-\d(?!\d))?)?[A-Z]?)(?!\d)",
    re.I,
)
# part of a multi-part document, right after the ID: "Part 10A", "_Part1A", ", VOLUME 1", "-Vol-2", "-ANNEX"
PART_AFTER_RE = re.compile(r"[\s,_-]*(?:Part[\s_]*(\d+)([A-Z])?(?![a-z])|Vol(?:ume)?\.?[\s_-]*(\d+)|(Annex)\b)", re.I)
# issue right after the ID: "Rev.2", "Rev. 1", "_Rev.1(15Nov...)", ", REVISION A", "w/Change 4", "-w-Change-4"
ISSUE_AFTER_RE = re.compile(
    r"[\s,_-]*(?:Rev(?:ision)?[.\s_-]*(\d|[A-Z](?![A-Za-z]))|(?:w[/-]?\s*|with[\s_-]+)change[\s_-]*(\d+))", re.I
)
# corrigendum after the issue (or the ID): " Corrigendum 1", "-Corr.1", "(1June2017)+Corrigendum1",
# "(15November2008)+Corr.1"; a corrected issue is its own issue ("Rev.2 Corr.1")
CORR_AFTER_RE = re.compile(r"[\s,_+-]*(?:\(\w+\)[\s_+-]*)?Corr(?:igendum)?[.\s_-]*(\d+)", re.I)
DASHES_RE = re.compile("[‐-―−]")  # PDFs often print the ID with typographic hyphens


def _ids(text: str) -> list[tuple[str, int]]:
    """(standard number with its part/volume, end offset) for each one named in `text`."""
    out = []
    for m in STD_ID_RE.finditer(text):
        doc_id, end = m.group(1).upper(), m.end()
        part = PART_AFTER_RE.match(text, end)
        suffix = ""
        if part and part.group(1):
            suffix = f" Part {part.group(1)}{(part.group(2) or '').upper()}"
        elif part and part.group(3):
            suffix = f" Vol {part.group(3)}"
        elif part and doc_id.startswith("NASA"):  # NASA issues annexes as documents; ECSS Word annexes belong to a file set
            suffix = " Annex"
        if suffix:
            doc_id, end = doc_id + suffix, part.end()
        out.append((doc_id, end))
    return out


def meta_from_text(text: str, prefer: str = "") -> tuple[str, str]:
    """(doc_id, issue) of the standard named in `text` (a cover page or a file name).

    The most frequent standard number wins (the page header repeats the document's own ID);
    `prefer` only breaks ties, then the first one named. Returns ("", "") when none is found.
    A part or volume of a multi-part document is part of its doc_id ("ECSS-E-HB-32-20 Part 1A").
    """
    text = DASHES_RE.sub("-", text)
    found = _ids(text)
    if not found:
        return "", ""
    ids = [i for i, _ in found]
    doc_id = max(ids, key=lambda i: (ids.count(i), i == prefer, -ids.index(i)))
    for i, end in found:
        if i != doc_id:
            continue
        m = ISSUE_AFTER_RE.match(text, end)
        issue = (f"Rev.{m.group(1).upper()}" if m.group(1) else f"w/Change {m.group(2)}") if m else ""
        corr = CORR_AFTER_RE.match(text, m.end() if m else end)
        if corr:
            issue = f"{issue} Corr.{corr.group(1)}".strip()
        if issue:
            return doc_id, issue
    return doc_id, ""


@lru_cache(maxsize=64)
def _manifest(folder: Path) -> dict[str, dict]:
    """file name -> record, from a download manifest in `folder` (a *.json with a "standards" list
    naming each `dest_filename`, as written by scripts/download_nasa_standards.py)."""
    out: dict[str, dict] = {}
    for m in sorted(folder.glob("*.json")):
        try:
            data = json.loads(m.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        for rec in data.get("standards", []) if isinstance(data, dict) else []:
            if isinstance(rec, dict) and rec.get("dest_filename"):
                out[rec["dest_filename"]] = rec
    return out


def manifest_record(path: Path) -> dict:
    return _manifest(path.parent).get(path.name, {})


def _cover_text(path: Path) -> tuple[str, str]:
    """(page header, first pages) of a document, where it names itself.

    A Word page header names the document itself, but appears only once in the file, so the
    frequency rule of `meta_from_text` cannot rely on it; it is returned separately.
    """
    try:
        if path.suffix.lower() == ".pdf":
            import pdfplumber

            with pdfplumber.open(str(path)) as pdf:
                return "", "\n".join(p.extract_text() or "" for p in pdf.pages[:2])
        if path.suffix.lower() == ".docx":
            import docx  # python-docx

            d = docx.Document(str(path))
            header = [p.text for s in d.sections[:1] for h in (s.first_page_header, s.header) for p in h.paragraphs]
            return "\n".join(header), "\n".join(p.text for p in d.paragraphs[:60])
    except Exception:  # unreadable here: fall back to the file name
        log.debug("cannot read cover of %s", path, exc_info=True)
    return "", ""


# ECSS PDF cover: "<ID>\n<date>\nSpace engineering\nTesting\nECSS Secretariat" (the series name may wrap)
COVER_TITLE_RE = re.compile(
    r"^Space\s+(?:engineering|product\s+assurance|project\s+management|sustainability|system)\s*\n(.+?)\n\s*ECSS Secretariat",
    re.M | re.S | re.I,
)


def title_from_cover(cover: str) -> str:
    m = COVER_TITLE_RE.search(cover)
    return " ".join(m.group(1).split()) if m and len(m.group(1)) < 300 else ""


def guess_meta(path: Path) -> dict:
    """doc_id and issue from the Word page header, else the cover page, else the file name, else a
    download manifest next to the file.

    The file name wins over the cover when it names a part of the same document ("…_Part10A",
    "…-VOLUME_1", "…-ANNEX") or a corrigendum of the cover's issue ("…_Rev.1+Corrigendum1") that the cover
    does not. `title` is added when the cover page (ECSS PDFs) or the manifest states one.
    """
    name_id, name_issue = meta_from_text(path.stem)
    header, cover = _cover_text(path)
    cover_id, cover_issue = meta_from_text(header, prefer=name_id)
    if not cover_id:
        cover_id, cover_issue = meta_from_text(cover, prefer=name_id)
    record = manifest_record(path)
    title = title_from_cover(cover) or " ".join(str(record.get("title", "")).split())
    if cover_id == name_id and " Corr." in name_issue and name_issue.startswith(cover_issue):
        cover_issue = name_issue  # the cover prints "Corrigendum 1" on a line of its own
    meta = {}
    if cover_id and name_id.startswith(cover_id + " "):
        meta = {"doc_id": name_id, "issue": cover_issue or name_issue}
    elif cover_id:
        meta = {"doc_id": cover_id, "issue": cover_issue or (name_issue if cover_id == name_id else "")}
    elif name_id:
        meta = {"doc_id": name_id, "issue": name_issue}
    elif record.get("doc_no"):
        meta = {"doc_id": str(record["doc_no"]).upper(), "issue": ""}
    if meta:
        return {**meta, "title": title} if title else meta
    match = re.match(r"^(.*?)[_ ]+(Rev\.?\s*\d+|Issue\s*\w+|[A-Z])$", path.stem)
    if match:
        return {"doc_id": match.group(1), "issue": match.group(2)}
    log.warning("No standard number found in %s; using %r. Add %s.meta.yaml to set doc_id and issue.",
                path.name, path.stem, path.stem)
    return {"doc_id": path.stem, "issue": ""}


def read_meta(path: Path) -> dict:
    """Document identity: sidecar `<name>.meta.yaml` fields override what is guessed."""
    sidecar = path.with_name(path.stem + ".meta.yaml")
    meta = (yaml.safe_load(sidecar.read_text(encoding="utf-8")) or {}) if sidecar.exists() else {}
    if "doc_id" in meta and "issue" in meta:
        return meta
    return {**guess_meta(path), **meta}


def _family(doc_id: str) -> str:
    d = doc_id.upper()
    if d.startswith("ECSS"):
        return "ecss"
    if d.startswith("NASA"):
        return "nasa"
    return "generic"


def _load_pdf(path: Path, doc: Document) -> None:
    from .ecss_parser import attach_tables, parse_ecss_lines
    from .nasa_parser import parse_nasa_lines
    from .pdf_extract import read_pdf

    content = read_pdf(path, doc.doc_id, doc.issue)
    doc.tables = content.tables
    family = _family(doc.doc_id)
    stats: dict = {}
    if family == "ecss":
        reqs, info = parse_ecss_lines(content.lines, doc.doc_id, doc.issue, handbook=is_handbook(doc.doc_id), table_lines=content.table_lines,
                                      stats=stats)
        attach_tables(doc.tables, info)
        doc.unparsed = stats["unparsed"]
    elif family == "nasa":
        reqs, doc.unparsed = parse_nasa_lines(content.lines, doc.doc_id, doc.issue,
                                              handbook=is_handbook(doc.doc_id), stats=stats)
        for t in doc.tables:  # clause = the requirement that points to the table
            owner = next((r for r in reqs if t.table_id in r.tables and r.kind == "requirement"), None)
            t.clause = owner.req_id if owner else ""
    else:
        reqs = parse_paragraphs([ln.text for ln in content.lines], doc.doc_id, doc.issue, is_handbook(doc.doc_id))
    doc.requirements = reqs
    doc.body_words, doc.unparsed_words = stats.get("body_words", 0), stats.get("unparsed_words", 0)


def _load_word_ecss(path: Path, doc: Document) -> None:
    from .ecss_parser import parse_ecss_lines, word_lines
    from .word_extract import word_to_text_lines

    stats: dict = {}
    reqs, _ = parse_ecss_lines(word_lines(word_to_text_lines(path)), doc.doc_id, doc.issue, recompute_letters=True,
                               handbook=is_handbook(doc.doc_id), stats=stats)
    doc.requirements, doc.unparsed = reqs, stats["unparsed"]
    doc.body_words, doc.unparsed_words = stats["body_words"], stats["unparsed_words"]


def _cache_key(path: Path, meta: dict | None = None) -> str:
    st = path.stat()
    sidecar = path.with_name(path.stem + ".meta.yaml")
    extra = sidecar.read_bytes() if sidecar.exists() else b""
    extra += json.dumps(meta, sort_keys=True).encode() if meta else b""
    h = hashlib.sha1(f"{_code_version()}|{path.name}|{st.st_size}|{st.st_mtime_ns}".encode() + extra)
    return h.hexdigest()[:16]


def load_document(path: Path, use_cache: bool = True, meta: dict | None = None) -> Document:
    """Parse one file. `meta` (doc_id, issue, title) skips identification: the catalogue passes the
    identity it gave the document's whole file group, so a Word copy cannot disagree with its PDF."""
    cache = path.parent / CACHE_DIR / f"{path.name}.{_cache_key(path, meta)}.pkl"
    if use_cache and cache.exists():
        try:
            return pickle.loads(cache.read_bytes())
        except Exception:  # stale or unreadable cache: re-parse
            pass

    meta = meta or read_meta(path)
    doc_id = str(meta.get("doc_id", path.stem))
    issue = str(meta.get("issue", ""))
    doc = Document(doc_id, issue, meta.get("title", ""), path.name)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        _load_pdf(path, doc)
    elif suffix in {".docx", ".doc"} and _family(doc_id) == "ecss":
        _load_word_ecss(path, doc)
    else:
        doc.requirements = parse_paragraphs(read_paragraphs(path), doc_id, issue, is_handbook(doc_id))

    if use_cache and suffix in {".pdf", ".doc", ".docx"}:
        try:
            cache.parent.mkdir(exist_ok=True)
            for old in cache.parent.glob(f"{path.name}.*.pkl"):
                old.unlink()
            cache.write_bytes(pickle.dumps(doc))
        except OSError:
            pass
    return doc


def load_folder(folder: Path, use_cache: bool = True) -> list[Document]:
    """The documents `corpus.yaml` selects, one source per doc_id (see catalogue.py)."""
    from .catalogue import build_catalogue

    return build_catalogue(folder, use_cache).load_documents(use_cache)
