"""Document identity from cover pages and file names. Standard numbers here are invented."""

import logging

import docx
import pytest

from spacestd.loaders import load_folder, meta_from_text, read_meta


@pytest.mark.parametrize("stem, expected", [
    ("ECSS-E-ST-99-01C_Rev.2", ("ECSS-E-ST-99-01C", "Rev.2")),
    ("ECSS-E-ST-99-01C_Rev.2(30July2021)", ("ECSS-E-ST-99-01C", "Rev.2")),
    ("ECSS-E-ST-99C_Rev.115November2008", ("ECSS-E-ST-99C", "Rev.1")),
    ("ECSS-E-ST-99-08C Rev.1", ("ECSS-E-ST-99-08C", "Rev.1")),
    ("ecss-e-st-99-10c_rev1", ("ECSS-E-ST-99-10C", "Rev.1")),
    ("ECSS-E-ST-99-02C", ("ECSS-E-ST-99-02C", "")),
    ("NASA-STD-9999A-w-Change-4-Revalidated-2025-09-05", ("NASA-STD-9999A", "w/Change 4")),
    ("NASA-STD-9998B", ("NASA-STD-9998B", "")),
    ("ECSS-E-ST-99-10C_Rev.2-Corr.1(1August2031)", ("ECSS-E-ST-99-10C", "Rev.2 Corr.1")),
    ("ECSS-Q-ST-99-26C-Rev.1(15March2031)+Corrigendum1(1June2031)", ("ECSS-Q-ST-99-26C", "Rev.1 Corr.1")),
    ("ECSS-Q-ST-99-14C Rev.1-Corrigendum1(2March2032)", ("ECSS-Q-ST-99-14C", "Rev.1 Corr.1")),
    ("ECSS-E-ST-99-12C(15November2028)+Corr.1(22February2031)", ("ECSS-E-ST-99-12C", "Corr.1")),
    ("random download", ("", "")),
])
def test_meta_from_file_name(stem, expected):
    assert meta_from_text(stem) == expected


def test_meta_from_cover_text():
    cover = (
        "ECSS‐E‐ST‐99‐10C Rev. 2\n12 March 2031\n"
        "Space engineering\nInvented structures standard\n"
        "This standard supersedes ECSS-E-ST-99-10C Rev.1 and refers to ECSS-E-ST-99-01C.\n"
        "ECSS-E-ST-99-10C Rev.2\n12 March 2031\n"  # page 2 header
    )
    assert meta_from_text(cover) == ("ECSS-E-ST-99-10C", "Rev.2")
    nasa = "NASA TECHNICAL STANDARD\nNASA-STD-9999A\nw/CHANGE 3\nSuperseding NASA-STD-9999\nNASA-STD-9999A"
    assert meta_from_text(nasa) == ("NASA-STD-9999A", "w/Change 3")
    corr = "ECSS-E-ST-99-10C Rev.2 Corrigendum 1\n1 August 2031\nprepared by the ECSS-E-ST-99-10C Rev.2 Working Group"
    assert meta_from_text(corr) == ("ECSS-E-ST-99-10C", "Rev.2 Corr.1")
    assert meta_from_text("ECSS-E-ST-99-10C Rev.2 Correction table") == ("ECSS-E-ST-99-10C", "Rev.2")


def _write_pdf(path, lines):
    """Smallest valid one-page PDF with the given text lines (no PDF library needed)."""
    ops = " ".join(f"BT /F1 12 Tf 72 {720 - 20 * i} Td ({t}) Tj ET" for i, t in enumerate(lines))
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(ops)} >>\nstream\n{ops}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = "%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    path.write_bytes(out.encode("latin-1"))


def test_pdf_cover_beats_file_name(tmp_path):
    path = tmp_path / "downloaded standard (1).pdf"
    _write_pdf(path, ["NASA-STD-9999A", "w/Change 2", "Invented cover page"])
    assert read_meta(path) == {"doc_id": "NASA-STD-9999A", "issue": "w/Change 2"}


def test_cover_frequency_beats_file_name(tmp_path):
    path = tmp_path / "ECSS-E-ST-99-01C.pdf"  # named after a standard the cover only refers to
    _write_pdf(path, ["ECSS-E-ST-99-10C Rev.2", "Invented page referring to ECSS-E-ST-99-01C", "ECSS-E-ST-99-10C Rev.2"])
    assert read_meta(path) == {"doc_id": "ECSS-E-ST-99-10C", "issue": "Rev.2"}
    assert meta_from_text("ECSS-E-ST-99-01C ECSS-E-ST-99-10C", prefer="ECSS-E-ST-99-10C")[0] == "ECSS-E-ST-99-10C"


def test_issue_from_name_when_cover_has_none(tmp_path):
    path = tmp_path / "ECSS-E-ST-99-01C_Rev.2(30July2031).pdf"
    _write_pdf(path, ["ECSS-E-ST-99-01C", "Invented cover page"])
    assert read_meta(path) == {"doc_id": "ECSS-E-ST-99-01C", "issue": "Rev.2"}


def test_docx_cover_and_header(tmp_path):
    path = tmp_path / "standard.docx"
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "ECSS-E-ST-99-08C Rev.1"
    d.add_paragraph("Invented title page")
    d.save(path)
    assert read_meta(path) == {"doc_id": "ECSS-E-ST-99-08C", "issue": "Rev.1"}


def test_docx_header_beats_references_on_title_page(tmp_path):
    path = tmp_path / "standard.docx"
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "ECSS-E-ST-99-08C Rev.1"  # printed once in a Word file
    d.add_paragraph("Invented foreword citing ECSS-S-ST-99-01 twice: ECSS-S-ST-99-01.")
    d.save(path)
    assert read_meta(path) == {"doc_id": "ECSS-E-ST-99-08C", "issue": "Rev.1"}


def test_sidecar_overrides_single_fields(tmp_path):
    path = tmp_path / "NASA-STD-9999A.pdf"
    _write_pdf(path, ["NASA-STD-9999A", "Invented cover page"])
    (tmp_path / "NASA-STD-9999A.meta.yaml").write_text("issue: w/Change 5\ntitle: Invented\n")
    assert read_meta(path) == {"doc_id": "NASA-STD-9999A", "issue": "w/Change 5", "title": "Invented"}


def test_unknown_name_warns(tmp_path, caplog):
    path = tmp_path / "scan.pdf"
    _write_pdf(path, ["Invented page without a standard number"])
    with caplog.at_level(logging.WARNING, logger="spacestd.loaders"):
        assert read_meta(path)["doc_id"] == "scan"
    assert "meta.yaml" in caplog.text


def test_pdf_preferred_over_word_with_download_names(tmp_path):
    _write_pdf(tmp_path / "ECSS-E-ST-99-01C_Rev.2(30July2031).pdf", ["ECSS-E-ST-99-01C Rev.2"])
    d = docx.Document()
    d.add_paragraph("ECSS-E-ST-99-01C Rev.2")
    d.save(tmp_path / "ECSS-E-ST-99-01C Rev.2 word.docx")
    [doc] = load_folder(tmp_path, use_cache=False)
    assert (doc.doc_id, doc.issue, doc.source_file) == ("ECSS-E-ST-99-01C", "Rev.2", "ECSS-E-ST-99-01C_Rev.2(30July2031).pdf")


def test_file_name_corrigendum_beats_cover_issue(tmp_path):
    path = tmp_path / "ECSS-Q-ST-99-26C-Rev.1(15March2031)+Corrigendum1(1June2031).pdf"
    _write_pdf(path, ["ECSS-Q-ST-99-26C Rev.1", "15 March 2031", "Corrigendum 1", "ECSS-Q-ST-99-26C Rev.1"])
    assert read_meta(path) == {"doc_id": "ECSS-Q-ST-99-26C", "issue": "Rev.1 Corr.1"}
    other = tmp_path / "ECSS-Q-ST-99-27C-Rev.2-Corr.1.pdf"  # the cover names a different issue: the cover wins
    _write_pdf(other, ["ECSS-Q-ST-99-27C Rev.3", "Invented cover", "ECSS-Q-ST-99-27C Rev.3"])
    assert read_meta(other) == {"doc_id": "ECSS-Q-ST-99-27C", "issue": "Rev.3"}
